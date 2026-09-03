"""Regression tests for tema / subtema / tono / grouping in app.py.

Streamlit is mocked so the module can be imported without a running server.
"""
import io
import re
import sys
import unittest
from unittest.mock import MagicMock, patch


class _Session(dict):
    def __getattr__(self, name):
        try:
            return self[name]
        except KeyError as exc:
            raise AttributeError(name) from exc

    def __setattr__(self, name, value):
        self[name] = value


def _install_streamlit_stub():
    st = MagicMock()
    st.session_state = _Session()
    st.secrets = {}
    st.set_page_config = lambda **_k: None
    sys.modules.setdefault("streamlit", st)
    return st


_install_streamlit_stub()

import app  # noqa: E402


MARCA = "Universidad Tecnológica de Bolívar"
ALIAS = "UTB"


class TestAnalisisOrder(unittest.TestCase):
    def test_usa_fragmentos_de_marca_primero(self):
        titulo = "El gobierno anuncia reforma tributaria nacional"
        resumen = "El Congreso debate impuestos. La UTB lanza una carrera de medicina deportiva en Cartagena."
        texto, hay = app._texto_clasificacion(titulo, resumen, MARCA, ALIAS)
        self.assertTrue(hay)
        self.assertIn("carrera", texto.lower())
        self.assertNotIn("reforma tributaria", texto.lower())

    def test_cae_a_resumen_si_no_hay_marca(self):
        titulo = "Corto"
        resumen = "El ministerio publicó un informe de movilidad urbana sobre el nuevo metro de Bogotá."
        texto, hay = app._texto_clasificacion(titulo, resumen, MARCA, ALIAS)
        self.assertFalse(hay)
        self.assertIn("movilidad", texto.lower())
        self.assertNotIn("Corto", texto)

    def test_cae_a_titulo_si_resumen_insuficiente(self):
        titulo = "Apertura del laboratorio de biotecnología marina"
        resumen = "N/A"
        texto, hay = app._texto_clasificacion(titulo, resumen, MARCA, ALIAS)
        self.assertFalse(hay)
        self.assertIn("laboratorio", texto.lower())


class TestEtiquetasEspecificas(unittest.TestCase):
    def test_no_emite_cubos_genericos(self):
        texto = (
            "La Universidad Tecnológica de Bolívar lanza una nueva carrera de medicina "
            "deportiva junto al hospital universitario de Cartagena."
        )
        sub = app._extraer_subtema_especifico(texto, MARCA, ALIAS)
        self.assertFalse(app._es_etiqueta_generica(sub), sub)
        self.assertNotIn(",", sub)
        self.assertNotEqual(sub.strip().lower(), "sin tema")
        self.assertNotIn("cobertura", sub.lower())
        self.assertGreaterEqual(len(sub.split()), 2)

    def test_etiqueta_generica_detecta_sintomas(self):
        for raw in (
            "Cobertura de información relevante",
            "Cobertura informativa general",
            "Sin tema",
            "Varios",
        ):
            self.assertTrue(app._es_etiqueta_generica(raw), raw)

    def test_investigacion_no_empieza_por_verbo(self):
        texto = (
            "La Universidad Tecnológica de Bolívar enfrenta una investigación "
            "por fallas operativas en sus laboratorios."
        )
        sub = app._extraer_subtema_especifico(texto, MARCA, ALIAS)
        self.assertFalse(app._es_verbo_cabeza(sub.split()[0]), sub)
        self.assertIn("investig", app.unidecode(sub.lower()))

    def test_lanzamiento_es_subtema_valido(self):
        self.assertTrue(app._validar_estructura_subtema("Lanzamiento de carrera deportiva"))

    def test_limpiar_tema_quita_comas(self):
        limpio = app.limpiar_tema("Educación superior, formación profesional")
        self.assertNotIn(",", limpio)
        self.assertTrue(limpio)
        self.assertFalse(app._es_etiqueta_generica(limpio))

    def test_fallback_del_clasificador_nunca_generico(self):
        clf = app.ClasificadorSubtema(MARCA, ALIAS)
        clf._last_blob = (
            "La UTB firma un convenio de formación profesional con el SENA "
            "para técnicos en logística portuaria."
        )
        et = clf._fallback(["UTB firma convenio con el SENA"])
        self.assertFalse(app._es_etiqueta_generica(et), et)
        self.assertNotIn(",", et)


class TestTonoSinMarca(unittest.TestCase):
    def test_sin_fragmentos_contexto_vacio(self):
        ctx = app.extraer_contexto_marca(
            "Crisis del sector avícola nacional",
            "Los productores piden ayudas al gobierno por el alza de insumos.",
            MARCA,
            ALIAS,
        )
        self.assertEqual(ctx, "")

    def test_con_fragmentos_contexto_no_vacio(self):
        ctx = app.extraer_contexto_marca(
            "UTB recibe premio de innovación",
            "La Universidad Tecnológica de Bolívar fue galardonada por su laboratorio.",
            MARCA,
            ALIAS,
        )
        self.assertTrue(ctx)
        self.assertTrue(app._menciona_marca_o_alias(ctx, MARCA, ALIAS))

    def test_snippet_pkl_usa_solo_contexto(self):
        snip = app._snippet_tono_pkl(
            "Titular negativo sobre el sector",
            "La UTB recibió un premio a la excelencia académica.",
        )
        self.assertIn("premio", snip.lower())
        self.assertNotIn("Titular negativo", snip)


class TestAgrupacion(unittest.TestCase):
    def test_misma_historia_comparte_etiquetas(self):
        titulos = [
            "UTB lanza carrera de medicina deportiva en Cartagena",
            "La UTB lanza carrera de medicina deportiva en Cartagena",
        ]
        resumenes = [
            "La Universidad Tecnológica de Bolívar presentó su nueva carrera de medicina deportiva.",
            "La Universidad Tecnológica de Bolívar presentó su nueva carrera de medicina deportiva.",
        ]
        temas, subtemas = app.etiquetar_sin_llm(titulos, resumenes, MARCA, ALIAS)
        self.assertEqual(subtemas[0], subtemas[1])
        self.assertEqual(temas[0], temas[1])
        self.assertFalse(app._es_etiqueta_generica(subtemas[0]), subtemas[0])
        self.assertNotIn(",", subtemas[0])
        self.assertNotIn(",", temas[0])

    def test_historias_distintas_no_comparten_subtema(self):
        titulos = [
            "UTB lanza carrera de medicina deportiva",
            "UTB es investigada por presuntas irregularidades en contrataciones",
        ]
        resumenes = [
            "La Universidad Tecnológica de Bolívar abre una carrera de medicina deportiva.",
            "La Universidad Tecnológica de Bolívar enfrenta una investigación por contrataciones.",
        ]
        _temas, subtemas = app.etiquetar_sin_llm(titulos, resumenes, MARCA, ALIAS)
        self.assertNotEqual(app.string_norm_label(subtemas[0]), app.string_norm_label(subtemas[1]))
        for s in subtemas:
            self.assertFalse(app._es_etiqueta_generica(s), s)

    def test_consistencia_propaga_solo_equivalentes(self):
        df = __import__("pandas").DataFrame({
            "Título": [
                "UTB lanza carrera de medicina deportiva",
                "La UTB lanza carrera de medicina deportiva",
                "Sancionan a otra universidad por fraude en becas",
            ],
            "Resumen - Aclaracion": [
                "La Universidad Tecnológica de Bolívar abre medicina deportiva.",
                "La Universidad Tecnológica de Bolívar abre medicina deportiva.",
                "Otra institución fue sancionada por fraude en el programa de becas.",
            ],
            "Tono IA": ["Positivo", "Neutro", "Negativo"],
            "Tema": ["Educación superior", "Educación superior", "Fraude en becas"],
            "Subtema": [
                "Lanzamiento de carrera deportiva",
                "Lanzamiento de carrera universitaria",
                "Sanción por fraude académico",
            ],
        })
        with patch.object(app, "get_embeddings_batch", return_value=[None, None, None]):
            out = app.aplicar_consistencia_grupos(
                df, "Título", "Resumen - Aclaracion",
                marca=MARCA, aliases=ALIAS,
            )
        self.assertEqual(out.loc[0, "Subtema"], out.loc[1, "Subtema"])
        self.assertEqual(out.loc[0, "Tono IA"], "Positivo")
        self.assertEqual(out.loc[1, "Tono IA"], "Positivo")
        self.assertNotEqual(out.loc[2, "Subtema"], out.loc[0, "Subtema"])
        self.assertEqual(out.loc[2, "Tono IA"], "Negativo")


class TestPklYHeuristica(unittest.TestCase):
    def _pipeline_temas(self, textos, clases):
        from sklearn.pipeline import make_pipeline
        from sklearn.feature_extraction.text import TfidfVectorizer
        from sklearn.naive_bayes import MultinomialNB
        import joblib

        pipe = make_pipeline(TfidfVectorizer(), MultinomialNB())
        pipe.fit(textos, clases)
        buf = io.BytesIO()
        joblib.dump(pipe, buf)
        buf.seek(0)
        return buf, list(dict.fromkeys(clases))

    def test_pkl_generico_se_reemplaza_dentro_del_vocabulario(self):
        buf, clases = self._pipeline_temas(
            [
                "lanzamiento carrera deportiva utb medicina",
                "cobertura general de noticias del dia",
                "convenio formacion profesional sena educacion",
            ],
            [
                "Educación superior",
                "Cobertura de información relevante",
                "Educación superior",
            ],
        )
        titulos = ["UTB lanza carrera de medicina deportiva"]
        resumenes = ["La Universidad Tecnológica de Bolívar abre medicina deportiva."]
        textos = [app._texto_clasificacion(titulos[0], resumenes[0], MARCA, ALIAS)[0]]
        pack = app.analizar_temas_con_pkl(textos, buf)
        self.assertIsNotNone(pack)
        temas, clases_out = pack
        self.assertEqual(len(temas), 1)
        self.assertIn(temas[0], clases_out)
        self.assertIn(temas[0], clases)
        self.assertFalse(app._es_etiqueta_generica(temas[0]), temas[0])

    def test_pkl_nunca_sale_del_vocabulario_ni_con_pred_generica(self):
        clases = [
            "Reconocimientos institucionales",
            "Gestión hospitalaria",
            "Cobertura de información relevante",
        ]
        buf, _ = self._pipeline_temas(
            [
                "reconocimiento premio distincion acr radiology",
                "hospital clinica pacientes internacion",
                "cobertura de informacion relevante noticias",
            ],
            clases,
        )
        texto = (
            "La Fundación Santa Fe de Bogotá fue reconocida por el American College "
            "of Radiology y se convirtió en la primera institución de Latinoamérica."
        )
        pack = app.analizar_temas_con_pkl([texto], buf)
        self.assertIsNotNone(pack)
        temas, clases_out = pack
        self.assertIn(temas[0], clases_out)
        forzado = app._resolver_etiqueta_pkl(
            "Cobertura de información relevante", texto, clases
        )
        self.assertIn(forzado, clases)
        self.assertFalse(app._es_etiqueta_generica(forzado), forzado)
        pd = __import__("pandas")
        df = pd.DataFrame({
            "Título": ["Distinción ACR"],
            "Resumen - Aclaracion": [texto],
            "Contexto analizado": [texto],
            "Tono IA": ["Positivo"],
            "Tema": ["Tema inventado que no está en el pkl"],
            "Subtema": ["Reconocimiento de ACR en Latinoamérica"],
        })
        with patch.object(app, "get_embeddings_batch", return_value=[None]):
            out = app.aplicar_consistencia_grupos(
                df, "Título", "Resumen - Aclaracion",
                marca="Fundación Santa Fe de Bogotá", aliases=None,
                vocabulario_tema=clases,
            )
        self.assertIn(out.loc[0, "Tema"], clases)

    def test_sin_pkl_etiquetas_especificas(self):
        titulos = ["Investigación por fallas operativas en el campus de la UTB"]
        resumenes = [
            "La Universidad Tecnológica de Bolívar enfrenta una investigación por fallas operativas en laboratorios."
        ]
        temas, subtemas = app.etiquetar_sin_llm(titulos, resumenes, MARCA, ALIAS)
        self.assertFalse(app._es_etiqueta_generica(subtemas[0]), subtemas[0])
        self.assertFalse(app._es_etiqueta_generica(temas[0]), temas[0])
        self.assertNotIn(",", subtemas[0])
        self.assertNotIn(",", temas[0])


CTX_FENAVI_AVICOLA = (
    "PRESIDENTE DE FENAVI El sector avícola realizará su encuentro anual "
    "para analizar las oportunidades de exportación, las perspectivas de "
    "crecimiento de la Industria y las estrategias ante la volatilidad de la tasa de cambio."
)
CTX_FLA_PERSONAS = (
    "Presidente - Fenavi Javier Díaz Presidente - Analdex Dirigente y "
    "exgerente - Fábrica de licores de Antioquia, FLA"
)
TITULO_FENAVI = "Presidente de Fenavi"


def _habla_de_fla(texto: str) -> bool:
    n = app.unidecode(str(texto or "").lower())
    return any(tok in n for tok in ("licor", "fla", "fabrica"))


class TestFenaviNoHeredaFla(unittest.TestCase):
    """Una nota del encuentro avícola no puede heredar FLA de otra nota del lote."""

    def test_nucleo_distinto_bloquea_fusión_aunque_corpus_grande(self):
        self.assertTrue(
            app._hechos_nucleo_distinto(CTX_FENAVI_AVICOLA, CTX_FLA_PERSONAS, "Fenavi", None)
        )
        self.assertFalse(
            app._pueden_compartir_subtema(
                CTX_FENAVI_AVICOLA, CTX_FLA_PERSONAS, "Fenavi", None, estricto=False
            )
        )
        self.assertFalse(
            app._historias_son_el_mismo_hecho(CTX_FENAVI_AVICOLA, CTX_FLA_PERSONAS)
        )

    def test_subtema_propio_no_menciona_fla(self):
        sub_a = app._extraer_subtema_especifico(CTX_FENAVI_AVICOLA, "Fenavi", None)
        self.assertFalse(app._es_etiqueta_generica(sub_a), sub_a)
        self.assertFalse(_habla_de_fla(sub_a), sub_a)
        blob = app.unidecode(sub_a.lower())
        self.assertTrue(
            any(tok in blob for tok in ("encuentro", "avicol", "export", "volatil", "cambio")),
            sub_a,
        )
        self.assertFalse(
            app._etiqueta_pertenece_al_texto(
                "Fábrica de licores de Antioquia", CTX_FENAVI_AVICOLA
            )
        )
        self.assertTrue(
            app._etiqueta_pertenece_al_texto(
                "Fábrica de licores de Antioquia", CTX_FLA_PERSONAS
            )
        )

    def test_etiquetar_sin_llm_no_copia_fla(self):
        titulos = [TITULO_FENAVI, TITULO_FENAVI]
        resumenes = [CTX_FENAVI_AVICOLA, CTX_FLA_PERSONAS]
        temas, subtemas = app.etiquetar_sin_llm(titulos, resumenes, "Fenavi", None)
        self.assertFalse(_habla_de_fla(subtemas[0]), subtemas[0])
        self.assertFalse(_habla_de_fla(temas[0]), temas[0])
        self.assertNotEqual(
            app.string_norm_label(subtemas[0]),
            app.string_norm_label(subtemas[1]),
        )
        self.assertFalse(app._es_etiqueta_generica(subtemas[0]), subtemas[0])
        self.assertFalse(app._es_etiqueta_generica(subtemas[1]), subtemas[1])
        self.assertNotIn(",", subtemas[0])
        self.assertNotIn(",", subtemas[1])

    def test_grupo_noticia_compartido_no_contamina(self):
        pd = __import__("pandas")
        df = pd.DataFrame({
            "Título": [TITULO_FENAVI, TITULO_FENAVI],
            "Resumen - Aclaracion": [CTX_FENAVI_AVICOLA, CTX_FLA_PERSONAS],
            "Contexto analizado": [CTX_FENAVI_AVICOLA, CTX_FLA_PERSONAS],
            "Grupo noticia": ["G00001", "G00001"],
            "Tono IA": ["Neutro", "Neutro"],
            "Tema": ["Fábrica de licores de Antioquia", "Fábrica de licores de Antioquia"],
            "Subtema": ["Fábrica de licores de Antioquia", "Fábrica de licores de Antioquia"],
        })
        with patch.object(app, "get_embeddings_batch", return_value=[None, None]):
            out = app.aplicar_consistencia_grupos(
                df, "Título", "Resumen - Aclaracion",
                marca="Fenavi", aliases=None,
            )
        self.assertFalse(_habla_de_fla(out.loc[0, "Subtema"]), out.loc[0, "Subtema"])
        self.assertFalse(_habla_de_fla(out.loc[0, "Tema"]), out.loc[0, "Tema"])
        self.assertNotEqual(
            app.string_norm_label(out.loc[0, "Subtema"]),
            app.string_norm_label("Fábrica de licores de Antioquia"),
        )

    def test_llm_no_reutiliza_fla_del_lote(self):
        pd = __import__("pandas")
        clf = app.ClasificadorSubtema("Fenavi", None)
        pbar = MagicMock()
        col = pd.Series([CTX_FENAVI_AVICOLA, CTX_FLA_PERSONAS])
        res = pd.Series([CTX_FENAVI_AVICOLA, CTX_FLA_PERSONAS])
        tit = pd.Series([TITULO_FENAVI, TITULO_FENAVI])
        with patch.object(app, "get_embeddings_batch", return_value=[None, None]), \
             patch.object(
                 clf, "_generar_etiqueta",
                 return_value="Fábrica de licores de Antioquia",
             ):
            subtemas = clf.procesar_lote(col, pbar, res, tit)
        self.assertEqual(len(subtemas), 2)
        self.assertFalse(_habla_de_fla(subtemas[0]), subtemas[0])
        self.assertFalse(app._es_etiqueta_generica(subtemas[0]), subtemas[0])
        self.assertNotIn(",", subtemas[0])
        self.assertNotEqual(
            app.string_norm_label(subtemas[0]),
            app.string_norm_label(subtemas[1]),
        )


class TestSubtemaCompleto(unittest.TestCase):
    """Subtemas completos para cualquier marca: no trocear nombres ni omitir el hecho."""

    def _assert_subtema_completo(self, sub, texto, marca):
        n = app.unidecode(sub.lower())
        self.assertFalse(app._es_etiqueta_generica(sub), sub)
        self.assertNotIn(",", sub)
        self.assertGreaterEqual(len(sub.split()), 2, sub)
        self.assertFalse(app._etiqueta_trocea_nombre(sub, texto, marca, None), sub)
        self.assertFalse(app._es_nombre_o_fragmento_marca(sub, marca, None), sub)

    def test_no_trocea_nombres_propios_de_marcas_distintas(self):
        casos = [
            (
                "Fundación Santa Fe de Bogotá",
                "La Fundación Santa Fe de Bogotá fue reconocida por el American College of Radiology "
                "y se convirtió en la primera institución de Latinoamérica en lograr esta distinción. "
                "La Fundación Santa Fe de Bogotá fue reconocida por el American College of Radiology (ACR) "
                "como ACR® International Center for Quality and Safety.",
                ("acr", "american college", "latinoameric", "quality", "distincion"),
            ),
            (
                "Hospital San Juan de Dios",
                "El Hospital San Juan de Dios fue reconocido por la World Health Organization (WHO) "
                "como centro de referencia en atención materna y se convirtió en el primer hospital "
                "de la región en lograr esta distinción.",
                ("who", "world health", "atencion", "materna", "distincion", "referencia"),
            ),
            (
                "Caja de Compensación Familiar de Fenalco",
                "La Caja de Compensación Familiar de Fenalco inauguró un centro de bienestar laboral "
                "en Cartagena para afiliados del sector comercio.",
                ("inaugur", "bienestar", "cartagena", "centro", "afiliad"),
            ),
        ]
        for marca, ctx, tokens_hecho in casos:
            with self.subTest(marca=marca):
                sub = app._extraer_subtema_especifico(ctx, marca, None)
                self._assert_subtema_completo(sub, ctx, marca)
                n = app.unidecode(sub.lower())
                self.assertTrue(any(tok in n for tok in tokens_hecho), sub)
                marca_n = app.unidecode(marca.lower())
                if "bogot" in n and "santa" in marca_n:
                    self.assertIn("santa fe de bogot", n, sub)
                if "dios" in n and "juan" in marca_n:
                    self.assertIn("san juan de dios", n, sub)

    def test_etiquetar_respeta_contexto_de_marca(self):
        marca = "Hospital San Juan de Dios"
        ctx = (
            "El Hospital San Juan de Dios fue reconocido por la World Health Organization (WHO) "
            "como centro de referencia en atención materna."
        )
        _temas, subtemas = app.etiquetar_sin_llm(
            ["Hospital recibe distinción internacional"],
            [ctx],
            marca,
            None,
        )
        self._assert_subtema_completo(subtemas[0], ctx, marca)
        n = app.unidecode(subtemas[0].lower())
        self.assertTrue(any(tok in n for tok in ("who", "world health", "atencion", "materna", "referencia")), subtemas[0])


CTX_ENCUENTRO_PANEL = (
    "El encuentro reunió al Dr. Giancarlo Buitrago, director de Investigaciones y Educación de LaCardio; "
    "al Dr. Gerardo Andrés Puentes Leal, líder de las Unidades de Endoscopia y del Servicio de Gastroenterología "
    "y jefe de Estudios y Epidemiología Clínicos del Hospital Serena del Mar; a Eddy Carolina Betancourt, "
    "directora científica de Alianza Team; y a Erick Eduardo Orozco Acosta, director del Doctorado en "
    "Inteligencia Artificial de la Universidad Simón Bolívar."
)

CTX_FORO_NUBARIA = (
    "El foro reunió a Marta Quilez, directora de Innovación Azul de Nubaria Tech; "
    "y a Pablo Reines, líder del Laboratorio de Robótica de la Universidad de Zelta."
)


class TestSubtemaFraseEvento(unittest.TestCase):
    """Subtema = frase de evento completa, no un n-grama ni un nombre de pila recortado."""

    def _assert_frase_evento(self, sub, texto):
        self.assertFalse(app._es_etiqueta_generica(sub), sub)
        self.assertNotIn(",", sub)
        self.assertGreaterEqual(len(sub.split()), 2, sub)
        self.assertFalse(app._subtema_de_baja_calidad(sub, texto), sub)
        n = app.unidecode(sub.lower())
        self.assertNotIn("reunio", n, sub)
        self.assertFalse(n.rstrip(".").endswith("giancarlo"), sub)
        self.assertFalse(n.rstrip(".").endswith("marta"), sub)
        self.assertFalse(n.rstrip(".").endswith("pablo"), sub)

    def test_encuentro_no_es_n_grama_ni_pila(self):
        sub = app._extraer_subtema_especifico(CTX_ENCUENTRO_PANEL, "LaCardio", ["lacardio"])
        self._assert_frase_evento(sub, CTX_ENCUENTRO_PANEL)
        n = app.unidecode(sub.lower())
        self.assertIn("encuentro", n)
        self.assertTrue(
            any(tok in n for tok in ("investig", "educacion", "especialist")),
            sub,
        )

    def test_encuentro_sin_marca_tambien_es_evento(self):
        sub = app._extraer_subtema_especifico(CTX_ENCUENTRO_PANEL, "", None)
        self._assert_frase_evento(sub, CTX_ENCUENTRO_PANEL)
        self.assertIn("encuentro", app.unidecode(sub.lower()))

    def test_rechaza_scrap_reunio_giancarlo(self):
        scrap = "Encuentro de reunio giancarlo"
        self.assertTrue(app._subtema_de_baja_calidad(scrap, CTX_ENCUENTRO_PANEL), scrap)
        regenerada = app._asegurar_etiqueta_especifica(
            scrap, CTX_ENCUENTRO_PANEL, "LaCardio", ["lacardio"]
        )
        self._assert_frase_evento(regenerada, CTX_ENCUENTRO_PANEL)
        self.assertNotEqual(app.string_norm_label(regenerada), app.string_norm_label(scrap))

    def test_foro_marca_ficticia_no_es_caso_especial(self):
        sub = app._extraer_subtema_especifico(CTX_FORO_NUBARIA, "Nubaria Tech", ["nubaria"])
        self._assert_frase_evento(sub, CTX_FORO_NUBARIA)
        n = app.unidecode(sub.lower())
        self.assertTrue(
            any(tok in n for tok in ("foro", "innovacion", "robotic", "laboratorio")),
            sub,
        )
        self.assertNotIn("reunio", n)


CTX_DIPORTO = (
    "Diporto propone una experiencia residencial de lujo que invita a descubrir "
    "una nueva forma de habitar e invertir en el Gran Canal de Serena del Mar. "
    "Único en su tipo, este desarrollo propone un entorno donde paisaje, "
    "arquitectura y vida se encuentran."
)
_COLLAGE_DIPORTO = "Canal de único en su tipo este desarrollo"


class TestSubtemaHechoNominal(unittest.TestCase):
    """El subtema es un encabezado gramatical del hecho, no un collage de keywords."""

    def _assert_residencial_lujo(self, sub):
        n = app.unidecode(sub.lower())
        self.assertFalse(re.search(r"canal de unico", n), sub)
        self.assertFalse(re.search(r"unico en su tipo este desarrollo", n), sub)
        self.assertFalse(app._es_etiqueta_generica(sub), sub)
        self.assertNotIn(",", sub)
        self.assertFalse(app._subtema_de_baja_calidad(sub, CTX_DIPORTO), sub)
        self.assertTrue(
            ("residencial" in n and ("lujo" in n or "serena" in n or "desarrollo" in n or "proyecto" in n))
            or ("desarrollo" in n and ("serena" in n or "lujo" in n or "residencial" in n))
            or ("proyecto" in n and ("residencial" in n or "lujo" in n)),
            sub,
        )

    def test_collage_es_baja_calidad(self):
        self.assertTrue(
            app._subtema_de_baja_calidad(_COLLAGE_DIPORTO, CTX_DIPORTO),
            _COLLAGE_DIPORTO,
        )

    def test_extraer_no_emite_collage(self):
        for marca, aliases in (
            ("Diporto", None),
            ("Serena del Mar", None),
            ("Diporto", ["Serena del Mar", "Gran Canal"]),
        ):
            with self.subTest(marca=marca):
                sub = app._extraer_subtema_especifico(CTX_DIPORTO, marca, aliases)
                self._assert_residencial_lujo(sub)

    def test_etiquetar_sin_llm_no_emite_collage(self):
        for marca, aliases in (
            ("Diporto", None),
            ("Serena del Mar", None),
        ):
            with self.subTest(marca=marca):
                _temas, subtemas = app.etiquetar_sin_llm(
                    ["Diporto lanza residencial de lujo"],
                    [CTX_DIPORTO],
                    marca,
                    aliases,
                )
                self._assert_residencial_lujo(subtemas[0])

    def test_no_une_tokens_sueltos_con_de(self):
        from pathlib import Path
        src = Path(app.__file__).with_name("calidad_etiquetas.py").read_text(encoding="utf-8")
        src += Path(app.__file__).read_text(encoding="utf-8")
        self.assertNotIn("join(top", src)
        self.assertNotIn("_palabras_contenido_evento", src)
        self.assertFalse(hasattr(app, "_palabras_contenido_evento"))


class TestCotaLLMYVelocidad(unittest.TestCase):
    """El etiquetado no puede cascada de refine ni colgarse en 1/N."""

    def test_generar_etiqueta_no_espiral_de_refine(self):
        import json as _json
        clf = app.ClasificadorSubtema(MARCA, ALIAS)
        calls = {"n": 0}

        def fake_create(*_a, **_k):
            calls["n"] += 1
            resp = MagicMock()
            choice = MagicMock()
            choice.message.content = _json.dumps({"subtema": "Varios"})
            resp.choices = [choice]
            resp.get = lambda k, d=None: {} if k == "usage" else d
            return resp

        blob = "Un texto demasiado vago para que la heurística cierre sola. Hechos varios del día."
        with patch.object(app.openai.ChatCompletion, "create", side_effect=fake_create):
            et = clf._generar_etiqueta(
                [blob],
                ["Noticia institucional del sector"],
                [blob],
            )
        self.assertLessEqual(calls["n"], app.MAX_LLM_CALLS_POR_ETIQUETA, calls["n"])
        self.assertLessEqual(getattr(clf, "_llm_calls_last", calls["n"]), app.MAX_LLM_CALLS_POR_ETIQUETA)
        self.assertTrue(str(et).strip())

    def test_heuristica_salta_llm_si_ya_es_alta_calidad(self):
        clf = app.ClasificadorSubtema(MARCA, ALIAS)
        calls = {"n": 0}

        def boom(*_a, **_k):
            calls["n"] += 1
            raise AssertionError("LLM no debía llamarse")

        blob = (
            "La Universidad Tecnológica de Bolívar lanza una nueva carrera de medicina "
            "deportiva junto al hospital universitario de Cartagena."
        )
        with patch.object(app.openai.ChatCompletion, "create", side_effect=boom):
            et = clf._generar_etiqueta(
                [blob],
                ["UTB lanza carrera de medicina deportiva"],
                [blob],
            )
        self.assertEqual(calls["n"], 0)
        self.assertFalse(app._es_etiqueta_generica(et), et)
        self.assertFalse(app._subtema_de_baja_calidad(et, blob), et)

    def test_validar_completa_sin_llm_por_defecto(self):
        calls = {"n": 0}

        def boom(*_a, **_k):
            calls["n"] += 1
            raise AssertionError("no LLM")

        with patch.object(app.openai.ChatCompletion, "create", side_effect=boom):
            et = app._validar_etiqueta_completa(
                "Lanzamiento de",
                titulos_grp=["UTB lanza carrera de medicina deportiva"],
                resumenes_grp=["La UTB abre medicina deportiva."],
                marca=MARCA, aliases=ALIAS, usar_llm=False,
            )
        self.assertEqual(calls["n"], 0)
        self.assertTrue(str(et).strip())

    def test_umbrales_corpus_grande_no_exigen_090(self):
        u = app._umbrales_adaptativos(313)
        self.assertLessEqual(u["sim_minima_agrupacion"], 0.84)
        self.assertLess(u["sim_minima_agrupacion"], 0.90)
        u_chico = app._umbrales_adaptativos(4)
        self.assertGreaterEqual(u_chico["sim_minima_agrupacion"], 0.90)

    def test_gpt5_nano_se_fuerza_a_default(self):
        env = {
            "OPENAI_CLASIF_MODEL": "gpt-5-nano-2025-08-07",
            "OPENAI_CLASIF_ALLOW_GPT5": "1",
        }
        with patch.dict(__import__("os").environ, env, clear=False):
            modelo = app._resolver_modelo_clasificacion()
        self.assertEqual(modelo, "gpt-4.1-nano-2025-04-14")
        self.assertEqual(modelo, app.MODELO_CLASIF_DEFAULT)
        self.assertTrue(app.advertencia_modelo_clasificacion())

    def test_80_singletons_como_mucho_4_chatcompletions(self):
        """80 grupos no pueden disparar 80 ChatCompletions: solo lotes de 25–40."""
        import json as _json
        n = 80
        clf = app.ClasificadorSubtema(MARCA, ALIAS)
        titulos = [f"Nota institucional {i} del dia" for i in range(n)]
        resumenes = [
            f"Cobertura general de hechos varios {i}. Informacion relevante del sector."
            for i in range(n)
        ]
        col = __import__("pandas").Series(resumenes)
        tit = __import__("pandas").Series(titulos)
        res = __import__("pandas").Series(resumenes)
        calls = {"n": 0, "models": []}
        progress = []

        def fake_create(*_a, **kw):
            calls["n"] += 1
            calls["models"].append(kw.get("model"))
            n_items = 30
            payload = {str(i): f"Hecho institucional {i}" for i in range(n_items)}
            resp = MagicMock()
            choice = MagicMock()
            choice.message.content = _json.dumps(payload)
            resp.choices = [choice]
            resp.get = lambda k, d=None: {} if k == "usage" else d
            return resp

        class _P:
            def progress(self, frac, text=""):
                progress.append(str(text))

        env = {
            "OPENAI_CLASIF_MODEL": "gpt-5-nano-2025-08-07",
            "GRILL_PULIR_SUBTEMAS": "1",
        }
        with patch.dict(__import__("os").environ, env, clear=False), \
             patch.object(app, "get_embeddings_batch", return_value=[None] * n), \
             patch.object(app, "_candidato_subtema_ok", return_value=False), \
             patch.object(app.openai.ChatCompletion, "create", side_effect=fake_create):
            app.refrescar_modelo_clasificacion()
            subtemas = clf.procesar_lote(col, _P(), res, tit)
        self.assertEqual(len(subtemas), n)
        self.assertLessEqual(calls["n"], 4, calls)
        self.assertTrue(calls["models"])
        self.assertTrue(all(m == "gpt-4.1-nano-2025-04-14" for m in calls["models"]), calls["models"])
        uno_a_uno = [
            t for t in progress
            if __import__("re").search(r"Etiquetando\s+\d+\s*/\s*\d+", t)
            and "lote" not in t.lower()
        ]
        self.assertFalse(uno_a_uno, uno_a_uno)
        self.assertTrue(
            any("lote" in t.lower() or "heurístic" in t.lower() or "heuristic" in t.lower()
                for t in progress),
            progress[:8],
        )

    def test_modelo_usado_es_41_nano_aunque_env_sea_gpt5(self):
        import json as _json
        seen = []

        def fake_create(*_a, **kw):
            seen.append(kw.get("model"))
            resp = MagicMock()
            choice = MagicMock()
            choice.message.content = _json.dumps({"0": "Lanzamiento de carrera deportiva"})
            resp.choices = [choice]
            resp.get = lambda k, d=None: {} if k == "usage" else d
            return resp

        env = {"OPENAI_CLASIF_MODEL": "gpt-5-nano-2025-08-07"}
        items = [{
            "idx": 0,
            "texto": "La UTB lanza una carrera de medicina deportiva en Cartagena.",
            "titulo": "UTB lanza carrera",
            "resumen": "La UTB abre medicina deportiva.",
            "heuristica": "Varios",
        }]
        with patch.dict(__import__("os").environ, env, clear=False), \
             patch.object(app.openai.ChatCompletion, "create", side_effect=fake_create):
            app.refrescar_modelo_clasificacion()
            app._pulir_subtemas_en_lotes(items, MARCA, ALIAS)
        self.assertEqual(seen, ["gpt-4.1-nano-2025-04-14"])


class TestContextoTonoFuente(unittest.TestCase):
    def test_contexto_combina_titulo_y_resumen(self):
        ctx = app.extraer_contexto_marca(
            "UTB inaugura laboratorio de biotecnología",
            "La Universidad Tecnológica de Bolívar abrió un laboratorio de biotecnología marina en Cartagena.",
            MARCA, ALIAS,
        )
        self.assertTrue(ctx)
        self.assertIn("laboratorio", ctx.lower())
        self.assertTrue(app._menciona_marca_o_alias(ctx, MARCA, ALIAS))
        self.assertGreaterEqual(len(ctx.split()), 6)

    def test_contexto_no_inventa_marca(self):
        ctx = app.extraer_contexto_marca(
            "Crisis del sector avícola nacional",
            "Los productores piden ayudas al gobierno por el alza de insumos.",
            MARCA, ALIAS,
        )
        self.assertEqual(ctx, "")

    def test_contexto_prioriza_titulo_resumen_sobre_cuerpo(self):
        cuerpo = ("HTML scrap sucio. " * 40) + "La UTB aparece al final del cuerpo sucio."
        ctx = app.extraer_contexto_marca(
            "UTB lanza carrera de medicina",
            "La UTB presenta medicina deportiva en Cartagena.",
            MARCA, ALIAS, cuerpo,
        )
        self.assertIn("medicina", ctx.lower())
        self.assertNotIn("HTML scrap", ctx)
        self.assertTrue(app._menciona_marca_o_alias(ctx, MARCA, ALIAS))

    def test_texto_clasificacion_sigue_orden_marca_resumen_titulo(self):
        texto, hay = app._texto_clasificacion(
            "El gobierno anuncia reforma tributaria nacional",
            "El Congreso debate impuestos. La UTB lanza una carrera de medicina deportiva en Cartagena.",
            MARCA, ALIAS,
        )
        self.assertTrue(hay)
        self.assertIn("carrera", texto.lower())
        self.assertNotIn("reforma tributaria", texto.lower())

    def test_tono_sin_mencion_es_neutro_sin_inventar(self):
        clf = app.ClasificadorTono(MARCA, ALIAS)
        self.assertFalse(clf._menciona_marca("Crisis del sector sin la institución"))
        det = app._tono_determinista(
            "Crisis del sector avícola nacional sin mención institucional",
            MARCA, ALIAS,
        )
        self.assertIsNone(det)

    def test_clasificar_core_sin_llm_expone_columnas(self):
        df = app.clasificar_noticias_core(
            ["UTB lanza carrera de medicina deportiva"],
            ["La Universidad Tecnológica de Bolívar abre medicina deportiva."],
            MARCA, ALIAS, usar_llm=False,
        )
        for col in ("Contexto analizado", "Tono IA", "Tema", "Subtema", "Grupo noticia"):
            self.assertIn(col, df.columns)
        self.assertFalse(app._es_etiqueta_generica(df.loc[0, "Subtema"]), df.loc[0, "Subtema"])

    def test_colab_txt_es_standalone_sin_import_app(self):
        from pathlib import Path
        src = Path(__file__).resolve().parent.joinpath("Grill_API_Colab.txt").read_text(encoding="utf-8")
        self.assertNotIn("import app", src)
        self.assertNotIn("from app", src)
        self.assertIn("clasificar_noticias_core", src)
        self.assertIn("gpt-4.1-nano-2025-04-14", src)
        compile(src, "Grill_API_Colab.txt", "exec")
        self.assertIn("detectar_duplicados_avanzado", src)
        self.assertIn("Duplicados", src)
        self.assertNotIn("api_key_opcional", src)

    def test_colab_tema_claro_y_upload_obvio(self):
        from pathlib import Path
        src = Path(__file__).resolve().parent.joinpath("Grill_API_Colab.txt").read_text(encoding="utf-8")
        self.assertIn("1. Sube el Excel (.xlsx) aquí", src)
        self.assertNotIn("#0d0d0d", src)
        self.assertNotIn("body_background_fill=\"#0d0d0d\"", src)
        self.assertTrue(
            "Soft" in src or "light" in src.lower() or "#f5f5f5" in src or "#ffffff" in src,
            "Colab debe usar tema claro",
        )


class TestVelocidadCorpusGrande(unittest.TestCase):
    """400 filas no pueden hacer n×n SequenceMatcher ni ChatCompletion de subtema/tema."""

    def test_grafo_equivalencia_bloqueado_menos_5000_pares(self):
        n = 400
        titulos = [f"Zeta{i} anuncia hecho puntual {i} en Cali" for i in range(n)]
        resumenes = [f"Resumen corto del hecho {i} en la jornada regional." for i in range(n)]
        app.construir_grafo_equivalencia(titulos, resumenes, marca="ZetaCorp")
        self.assertLess(app._PARES_GRAFO_REVISADOS, 5000, app._PARES_GRAFO_REVISADOS)
        self.assertLess(app._PARES_GRAFO_REVISADOS, n * 20)

    def test_consistencia_400_titulos_menos_2s(self):
        import time as _time
        pd = __import__("pandas")
        n = 400
        titulos = [f"Zeta{i} anuncia hecho puntual {i} en Cali" for i in range(n)]
        resumenes = [f"Resumen corto del hecho {i} en la jornada regional." for i in range(n)]
        df = pd.DataFrame({
            "Título": titulos,
            "Resumen - Aclaracion": resumenes,
            "Tono IA": ["Neutro"] * n,
            "Tema": ["Hecho puntual"] * n,
            "Subtema": [f"Hecho puntual {i}" for i in range(n)],
        })
        t0 = _time.perf_counter()
        with patch.object(app, "get_embeddings_batch", return_value=[None] * n):
            out = app.aplicar_consistencia_grupos(
                df, "Título", "Resumen - Aclaracion", marca="ZetaCorp",
            )
        elapsed = _time.perf_counter() - t0
        self.assertEqual(len(out), n)
        self.assertLess(elapsed, 2.0, elapsed)
        self.assertLess(app._PARES_GRAFO_REVISADOS, 5000, app._PARES_GRAFO_REVISADOS)
        self.assertTrue(all(str(g).startswith("G") for g in out["Grupo noticia"]))

    def test_clasificar_core_default_cero_chatcompletions(self):
        n = 80
        titulos = [
            f"Zeta{i} inaugura laboratorio de biotecnología marina en Cartagena"
            for i in range(n)
        ]
        resumenes = [
            f"La empresa Zeta{i} abre un laboratorio de biotecnología marina en Cartagena."
            for i in range(n)
        ]
        calls = {"n": 0}

        def boom(*_a, **_k):
            calls["n"] += 1
            raise AssertionError("ChatCompletion no debe llamarse con flags por defecto")

        with patch.object(app, "get_embeddings_batch", return_value=[None] * n), \
             patch.object(app.openai.ChatCompletion, "create", side_effect=boom), \
             patch.dict(__import__("os").environ, {
                 "GRILL_PULIR_SUBTEMAS": "0",
                 "GRILL_PULIR_TEMAS": "0",
                 "GRILL_PULIR_TONO": "0",
             }, clear=False):
            df = app.clasificar_noticias_core(titulos, resumenes, "ZetaCorp", usar_llm=True)
        self.assertEqual(calls["n"], 0)
        self.assertEqual(len(df), n)
        self.assertIn("Grupo noticia", df.columns)
        self.assertIn("Subtema", df.columns)
        self.assertTrue(all(not app._es_etiqueta_generica(s) for s in df["Subtema"].tolist()[:5]))


KM_DUP = {
    "idnoticia": "ID Noticia",
    "hora": "Hora",
    "medio": "Medio",
    "tipodemedio": "Tipo de Medio",
    "titulo": "Título",
    "link_nota": "Link Nota",
    "link_streaming": "Link (Streaming - Imagen)",
    "menciones": "Menciones - Empresa",
    "idduplicada": "ID duplicada",
}


class TestDuplicadosLimpiezaGrill(unittest.TestCase):
    """Duplicada = misma publicación (limpieza_grill). Título similar NO es duplicado."""

    def test_titulos_similares_urls_distintas_no_son_duplicado(self):
        rows = [
            {
                "ID Noticia": "100",
                "Hora": "08:00",
                "Medio": "El Tiempo",
                "Tipo de Medio": "Internet",
                "Título": "UTB lanza carrera de medicina deportiva en Cartagena",
                "Link Nota": {"url": "https://eltiempo.com/nota-a"},
                "Link (Streaming - Imagen)": None,
                "Menciones - Empresa": "UTB",
                "is_duplicate": False,
            },
            {
                "ID Noticia": "101",
                "Hora": "09:00",
                "Medio": "El Tiempo",
                "Tipo de Medio": "Internet",
                "Título": "La UTB lanza carrera de medicina deportiva en Cartagena",
                "Link Nota": {"url": "https://eltiempo.com/nota-b"},
                "Link (Streaming - Imagen)": None,
                "Menciones - Empresa": "UTB",
                "is_duplicate": False,
            },
        ]
        out = app.detectar_duplicados_avanzado(rows, KM_DUP)
        self.assertFalse(out[0]["is_duplicate"])
        self.assertFalse(out[1]["is_duplicate"])

    def test_misma_url_y_mencion_es_duplicada(self):
        rows = [
            {
                "ID Noticia": "200",
                "Hora": "",
                "Medio": "Semana",
                "Tipo de Medio": "Internet",
                "Título": "Titular A",
                "Link Nota": {"url": "https://www.semana.com/nota-x/"},
                "Link (Streaming - Imagen)": None,
                "Menciones - Empresa": "UTB",
                "is_duplicate": False,
            },
            {
                "ID Noticia": "201",
                "Hora": "",
                "Medio": "Semana",
                "Tipo de Medio": "Internet",
                "Título": "Titular distinto por completo",
                "Link Nota": {"url": "http://semana.com/nota-x"},
                "Link (Streaming - Imagen)": None,
                "Menciones - Empresa": "UTB",
                "is_duplicate": False,
            },
        ]
        out = app.detectar_duplicados_avanzado(rows, KM_DUP)
        self.assertFalse(out[0]["is_duplicate"])
        self.assertTrue(out[1]["is_duplicate"])
        self.assertEqual(str(out[1]["ID duplicada"]), "200")

    def test_radio_mismo_slot_es_duplicado_hora_distinta_no(self):
        base = {
            "Tipo de Medio": "Radio",
            "Medio": "Caracol Radio",
            "Menciones - Empresa": "UTB",
            "Link Nota": None,
            "Link (Streaming - Imagen)": None,
            "is_duplicate": False,
            "Título": "Noticiero de la mañana",
        }
        rows = [
            {**base, "ID Noticia": "301", "Hora": "06:00"},
            {**base, "ID Noticia": "302", "Hora": "06:00"},
            {**base, "ID Noticia": "303", "Hora": "07:00"},
        ]
        out = app.detectar_duplicados_avanzado(rows, KM_DUP)
        self.assertFalse(out[0]["is_duplicate"])
        self.assertTrue(out[1]["is_duplicate"])
        self.assertEqual(str(out[1]["ID duplicada"]), "301")
        self.assertFalse(out[2]["is_duplicate"])

    def test_detectar_duplicados_no_usa_sequence_matcher_de_titulo(self):
        import inspect
        src = inspect.getsource(app.detectar_duplicados_avanzado)
        # Title-as-duplicate (internet buckets) stays gone. SequenceMatcher is
        # only allowed inside the radio/TV fecha+hora slot.
        self.assertNotIn("tb[(medio", src)
        internet = src.split("elif tipo in")[0]
        self.assertNotIn("SequenceMatcher", internet)


class TestGrupoNoticiaCoberturaSimilar(unittest.TestCase):
    """Grupo noticia = cobertura similar; copia tono/tema/subtema del representante."""

    def test_titulos_similares_urls_distintas_mismo_grupo_y_etiquetas(self):
        pd = __import__("pandas")
        df = pd.DataFrame({
            "Título": [
                "UTB lanza carrera de medicina deportiva en Cartagena",
                "La UTB lanza carrera de medicina deportiva en Cartagena",
            ],
            "Resumen - Aclaracion": [
                "La Universidad Tecnológica de Bolívar presentó su nueva carrera de medicina deportiva.",
                "La Universidad Tecnológica de Bolívar presentó su nueva carrera de medicina deportiva.",
            ],
            "Tono IA": ["Positivo", "Neutro"],
            "Tema": ["Educación superior", "Formación profesional"],
            "Subtema": [
                "Lanzamiento de carrera deportiva",
                "Apertura de programa académico",
            ],
        })
        with patch.object(app, "get_embeddings_batch", return_value=[None, None]):
            out = app.aplicar_consistencia_grupos(
                df, "Título", "Resumen - Aclaracion",
                marca=MARCA, aliases=ALIAS,
            )
        self.assertEqual(out.loc[0, "Grupo noticia"], out.loc[1, "Grupo noticia"])
        self.assertTrue(str(out.loc[0, "Grupo noticia"]).startswith("G"))
        self.assertEqual(out.loc[0, "Tono IA"], out.loc[1, "Tono IA"])
        self.assertEqual(out.loc[0, "Tema"], out.loc[1, "Tema"])
        self.assertEqual(out.loc[0, "Subtema"], out.loc[1, "Subtema"])

    def test_resumenes_similares_agrupan_aunque_titulos_diferan(self):
        pd = __import__("pandas")
        resumen = (
            "La Universidad Tecnológica de Bolívar inauguró el laboratorio de "
            "biotecnología marina en el campus de Cartagena con apoyo del ministerio."
        )
        df = pd.DataFrame({
            "Título": [
                "UTB abre laboratorio marino",
                "Nueva sede científica en Cartagena",
            ],
            "Resumen - Aclaracion": [resumen, resumen],
            "Tono IA": ["Positivo", "Neutro"],
            "Tema": ["Investigación científica", "Campus universitario"],
            "Subtema": [
                "Inauguración de laboratorio marino",
                "Apertura de campus costero",
            ],
        })
        with patch.object(app, "get_embeddings_batch", return_value=[None, None]):
            out = app.aplicar_consistencia_grupos(
                df, "Título", "Resumen - Aclaracion",
                marca=MARCA, aliases=ALIAS,
            )
        self.assertEqual(out.loc[0, "Grupo noticia"], out.loc[1, "Grupo noticia"])
        self.assertEqual(out.loc[0, "Tono IA"], out.loc[1, "Tono IA"])
        self.assertEqual(out.loc[0, "Tema"], out.loc[1, "Tema"])
        self.assertEqual(out.loc[0, "Subtema"], out.loc[1, "Subtema"])

    def test_prefijo_de_70_titulares_iguales_si_agrupa(self):
        """PR #13 saltaba cubetas >64; 70 copias del mismo titular deben ser un grupo."""
        pd = __import__("pandas")
        n = 70
        tit = "UTB lanza carrera de medicina deportiva en Cartagena"
        res = "La Universidad Tecnológica de Bolívar presentó medicina deportiva."
        df = pd.DataFrame({
            "Título": [tit] * n,
            "Resumen - Aclaracion": [res] * n,
            "Tono IA": ["Positivo"] + ["Neutro"] * (n - 1),
            "Tema": ["Educación superior"] * n,
            "Subtema": ["Lanzamiento de carrera deportiva"] * n,
        })
        with patch.object(app, "get_embeddings_batch", return_value=[None] * n):
            out = app.aplicar_consistencia_grupos(
                df, "Título", "Resumen - Aclaracion",
                marca=MARCA, aliases=ALIAS,
            )
        self.assertEqual(len(set(out["Grupo noticia"])), 1)
        self.assertTrue(all(t == "Positivo" for t in out["Tono IA"]))

    def test_columnas_xlsx_obligatorias_en_orden(self):
        esperadas = [
            "ID Noticia", "Fecha", "Hora", "Medio", "Tipo de Medio",
            "Sección - Programa", "Región", "Título", "Tono IA", "Tema", "Subtema",
            "Link Nota", "Resumen - Aclaracion", "Link (Streaming - Imagen)",
            "Menciones - Empresa", "ID duplicada", "Cuerpo Completo",
            "Contexto analizado", "Coincidencia marca", "Origen coincidencia",
            "Tono", "Grupo noticia",
        ]
        self.assertEqual(list(app.COLUMNAS_XLSX_OBLIGATORIAS), esperadas)
        orden = app.columnas_salida_xlsx()
        self.assertEqual(orden[: len(esperadas)], esperadas)

    def test_n400_pares_bloqueados_no_nxn(self):
        n = 400
        titulos = [f"Zeta{i} anuncia hecho puntual {i} en Cali" for i in range(n)]
        resumenes = [f"Resumen corto del hecho {i} en la jornada regional." for i in range(n)]
        app.construir_grafo_equivalencia(titulos, resumenes, marca="ZetaCorp")
        nxn = n * (n - 1) // 2
        self.assertLess(app._PARES_GRAFO_REVISADOS, 5000)
        self.assertLess(app._PARES_GRAFO_REVISADOS, nxn // 4)


CTX_UAO_BYLINE = (
    "Estudiante en formación Universidad Autónoma de Occidente "
    "Editor web y periodista egresado de la Universidad Autónoma de Occidente. "
    "Editora web y periodista egresada de la Universidad Autónoma de Occidente."
)
MARCA_UAO = "Universidad Autónoma de Occidente"
ALIAS_UAO = "UAO"

CASOS_SCRAP_UAO = [
    ("¿Qué clima hará en Cali este lunes?", "Clima de hara"),
    ("¿La Tierra es redonda?", "Tierra de redonda"),
    ("¿Cómo saber si tengo una multa de tránsito en Cali?", "Tengo de multa"),
    ("¿Dónde jugará América de Cali este domingo?", "Jugara de america"),
    ("¿Qué hacer si una persona necesita ayuda en una emergencia?", "Necesita de ayuda"),
]


class TestScrapsYAutoriaBylines(unittest.TestCase):
    """Reglas del dossier UAO 410: nada de 'X de Y' pegado; bylines ≠ el hecho."""

    def test_cinco_scraps_no_salen_con_contexto_byline(self):
        for titulo, scrap in CASOS_SCRAP_UAO:
            with self.subTest(titulo=titulo):
                temas, subs = app.etiquetar_sin_llm(
                    [titulo], [CTX_UAO_BYLINE], MARCA_UAO, ALIAS_UAO
                )
                self.assertNotEqual(
                    app.string_norm_label(subs[0]), app.string_norm_label(scrap), subs[0]
                )
                self.assertTrue(app._es_pegamento_de_tokens(scrap, titulo), scrap)
                self.assertFalse(app._es_pegamento_de_tokens(subs[0], titulo + " " + CTX_UAO_BYLINE), subs[0])
                self.assertTrue(app._es_subtema_autoria(subs[0]), subs[0])
                self.assertEqual(
                    app.string_norm_label(temas[0]),
                    app.string_norm_label("Estudiantes y egresados"),
                )

    def test_byline_marca_no_universitaria(self):
        marca = "Hospital Nubaria"
        ctx = (
            "Editora web y periodista de Hospital Nubaria. "
            "Comunicador del Hospital Nubaria."
        )
        titulo = "¿Qué clima hará en Cali este lunes?"
        temas, subs = app.etiquetar_sin_llm([titulo], [ctx], marca, None)
        self.assertTrue(app._es_subtema_autoria(subs[0]), subs[0])
        self.assertNotEqual(app.string_norm_label(subs[0]), app.string_norm_label("Clima de hara"))
        n = app.string_norm_label(temas[0])
        self.assertTrue(
            any(tok in n for tok in ("autoria", "periodist", "egresad")),
            temas[0],
        )

    def test_evento_real_de_la_marca_no_es_articulo(self):
        ctx = (
            "La Universidad Autónoma de Occidente recibió la acreditación de alta calidad "
            "del Ministerio de Educación Nacional y se convirtió en referente del Valle."
        )
        titulo = "La UAO recibe acreditación de alta calidad"
        temas, subs = app.etiquetar_sin_llm([titulo], [ctx], MARCA_UAO, ALIAS_UAO)
        self.assertFalse(app._es_subtema_autoria(subs[0]), subs[0])
        n = app.unidecode(subs[0].lower())
        self.assertTrue(
            any(tok in n for tok in ("acredit", "alta calidad", "certific", "ministerio")),
            subs[0],
        )
        self.assertFalse(app._es_pegamento_de_tokens(subs[0], ctx), subs[0])

    def test_pkl_autoria_mapea_a_egresados(self):
        helper = TestPklYHeuristica()
        buf, clases = helper._pipeline_temas(
            [
                "egresados alumni graduados universidad periodista",
                "acreditacion alta calidad institucional ministerio",
                "investigacion laboratorios campus",
            ],
            ["egresados", "Acreditación institucional", "Investigación"],
        )
        with patch.object(app, "get_embeddings_batch", return_value=[None]), \
             patch.object(app.openai.ChatCompletion, "create", side_effect=AssertionError("no LLM")):
            out = app.clasificar_noticias_core(
                ["¿Qué clima hará en Cali este lunes?"],
                [CTX_UAO_BYLINE],
                MARCA_UAO, ALIAS_UAO,
                pkl_tema=buf, usar_llm=False,
            )
        tema = out.loc[0, "Tema"]
        self.assertIn(tema, clases)
        self.assertEqual(app.string_norm_label(tema), app.string_norm_label("egresados"), tema)
        self.assertTrue(app._es_subtema_autoria(out.loc[0, "Subtema"]), out.loc[0, "Subtema"])

    def test_preguntas_sin_marca_no_pegan_tokens_con_de(self):
        for titulo, scrap in CASOS_SCRAP_UAO:
            with self.subTest(titulo=titulo):
                sub = app._extraer_subtema_especifico(titulo, "", None)
                self.assertNotEqual(app.string_norm_label(sub), app.string_norm_label(scrap), sub)
                self.assertFalse(app._es_pegamento_de_tokens(sub, titulo), sub)
                self.assertFalse(app._es_subtema_autoria(sub), sub)

    def test_tono_byline_sigue_siendo_neutro_con_mencion(self):
        self.assertTrue(app._menciona_marca_o_alias(CTX_UAO_BYLINE, MARCA_UAO, ALIAS_UAO))
        self.assertTrue(app._es_contexto_solo_autoria(
            "¿Qué clima hará en Cali este lunes?", CTX_UAO_BYLINE, MARCA_UAO, ALIAS_UAO
        ))
        clf = app.ClasificadorTono(MARCA_UAO, ALIAS_UAO)
        self.assertTrue(clf._menciona_marca(CTX_UAO_BYLINE))


class TestVelocidadDossier410(unittest.TestCase):
    """410 notas no pueden tardar ~483s: heurística + 0 ChatCompletions por defecto."""

    def test_410_byline_sin_llm_muy_por_debajo_de_2_min(self):
        import time as _time
        n = 410
        titulos = [CASOS_SCRAP_UAO[i % len(CASOS_SCRAP_UAO)][0] + f" ({i})" for i in range(n)]
        resumenes = [CTX_UAO_BYLINE] * n
        t0 = _time.perf_counter()
        with patch.object(app, "get_embeddings_batch", return_value=[None] * n), \
             patch.object(app.openai.ChatCompletion, "create", side_effect=AssertionError("no LLM")):
            df = app.clasificar_noticias_core(
                titulos, resumenes, MARCA_UAO, ALIAS_UAO, usar_llm=False
            )
        elapsed = _time.perf_counter() - t0
        self.assertEqual(len(df), n)
        self.assertLess(elapsed, 30.0, elapsed)
        scraps = {app.string_norm_label(s) for _, s in CASOS_SCRAP_UAO}
        produced = {app.string_norm_label(s) for s in df["Subtema"].tolist()}
        self.assertFalse(produced & scraps, produced & scraps)
        self.assertTrue(all(app._es_subtema_autoria(s) for s in df["Subtema"].head(15)), df["Subtema"].head(3).tolist())
        self.assertTrue(
            all(app.string_norm_label(t) == app.string_norm_label("Estudiantes y egresados") for t in df["Tema"].head(15)),
            df["Tema"].head(3).tolist(),
        )
        self.assertTrue(app._LAST_PHASE_TIMINGS, "faltan timings de fase")


TITULO_UAO_ACRED_A = (
    "Universidad Autónoma de Occidente recibe máxima acreditación del "
    "Ministerio de Educación por su calidad académica e impacto regional"
)
TITULO_UAO_ACRED_B = (
    "Universidad Autónoma de Occidente renueva su acreditación de alta calidad por 8 años"
)
RESUMEN_UAO_ACRED_A = (
    "La Universidad Autónoma de Occidente recibió la acreditación de alta calidad "
    "del Ministerio de Educación Nacional por su calidad académica e impacto regional."
)
RESUMEN_UAO_ACRED_B = (
    "La Universidad Autónoma de Occidente renovó su acreditación de alta calidad por 8 años."
)


class TestCollageDeRechazado(unittest.TestCase):
    """Bug 1: keyword collage glued with 'de' must be rejected and repaired."""

    def test_collage_x_de_y_se_rechaza(self):
        scraps = (
            "Calidad de académica",
            "Canal de único",
            "Clima de hara",
            "Tierra de redonda",
            "Tengo de multa",
        )
        fuente = (
            "Diporto propone una experiencia residencial de lujo en el Gran Canal. "
            "Único en su tipo, este desarrollo invita a habitar."
        )
        for scrap in scraps:
            with self.subTest(scrap=scrap):
                self.assertTrue(app._es_pegamento_de_tokens(scrap, fuente), scrap)
                self.assertTrue(app._subtema_de_baja_calidad(scrap, fuente), scrap)
                repaired = app._asegurar_etiqueta_especifica(scrap, fuente, "Diporto", None)
                self.assertFalse(app._es_pegamento_de_tokens(repaired, fuente), repaired)
                self.assertGreaterEqual(len(repaired.split()), 4, repaired)
                self.assertTrue(app._validar_estructura_subtema(repaired), repaired)

    def test_frase_gramatical_4_palabras_se_usa(self):
        fuente = CTX_DIPORTO
        sub = app._extraer_subtema_especifico(fuente, "Diporto", None)
        self.assertGreaterEqual(len(sub.split()), 4, sub)
        self.assertFalse(app._es_pegamento_de_tokens(sub, fuente), sub)
        self.assertIn("residencial", app.unidecode(sub.lower()))


class TestUaoAcreditacionUnifica(unittest.TestCase):
    """Bug 2: same accreditation fact → one subtema and one Grupo noticia."""

    def test_dos_titulos_misma_acreditacion(self):
        pd = __import__("pandas")
        df = pd.DataFrame({
            "Título": [TITULO_UAO_ACRED_A, TITULO_UAO_ACRED_B],
            "Resumen - Aclaracion": [RESUMEN_UAO_ACRED_A, RESUMEN_UAO_ACRED_B],
            "Contexto analizado": [RESUMEN_UAO_ACRED_A, RESUMEN_UAO_ACRED_B],
            "Tono IA": ["Positivo", "Neutro"],
            "Tema": ["Reconocimiento de alta calidad académica", "Reconocimiento de calidad institucional"],
            "Subtema": [
                "Reconocimiento de alta calidad académica",
                "Reconocimiento de calidad institucional",
            ],
        })
        with patch.object(app, "get_embeddings_batch", return_value=[None, None]):
            temas, subs = app.etiquetar_sin_llm(
                [TITULO_UAO_ACRED_A, TITULO_UAO_ACRED_B],
                [RESUMEN_UAO_ACRED_A, RESUMEN_UAO_ACRED_B],
                MARCA_UAO, ALIAS_UAO,
            )
            out = app.aplicar_consistencia_grupos(
                df, "Título", "Resumen - Aclaracion",
                marca=MARCA_UAO, aliases=ALIAS_UAO,
            )
        self.assertEqual(app.string_norm_label(subs[0]), app.string_norm_label(subs[1]), subs)
        self.assertGreaterEqual(len(subs[0].split()), 4, subs[0])
        self.assertTrue(
            any(tok in app.unidecode(subs[0].lower()) for tok in ("acredit", "calidad")),
            subs[0],
        )
        self.assertFalse(app._es_pegamento_de_tokens(subs[0], RESUMEN_UAO_ACRED_A), subs[0])
        self.assertEqual(out.loc[0, "Grupo noticia"], out.loc[1, "Grupo noticia"])
        self.assertEqual(out.loc[0, "Subtema"], out.loc[1, "Subtema"])
        self.assertEqual(out.loc[0, "Tema"], out.loc[1, "Tema"])
        self.assertEqual(temas[0], temas[1])


class TestAfiliacionGenerica(unittest.TestCase):
    """Affiliation/byline of any client → Neutro + Estudiantes y egresados."""

    def test_egresado_estudiante_byline_neutro_y_tema(self):
        with patch.object(app, "get_embeddings_batch", return_value=[None]), \
             patch.object(app.openai.ChatCompletion, "create", side_effect=AssertionError("no LLM")):
            out = app.clasificar_noticias_core(
                ["¿Qué clima hará en Cali este lunes?"],
                [CTX_UAO_BYLINE],
                MARCA_UAO, ALIAS_UAO, usar_llm=False,
            )
        self.assertEqual(out.loc[0, "Tono IA"], "Neutro")
        self.assertEqual(
            app.string_norm_label(out.loc[0, "Tema"]),
            app.string_norm_label("Estudiantes y egresados"),
        )
        self.assertTrue(app._es_subtema_autoria(out.loc[0, "Subtema"]), out.loc[0, "Subtema"])
        self.assertGreaterEqual(len(str(out.loc[0, "Subtema"]).split()), 4)
        self.assertNotIn("acredit", app.unidecode(str(out.loc[0, "Subtema"]).lower()))

    def test_afiliacion_es_generica_no_solo_uao(self):
        marca = "Hospital Nubaria"
        ctx = (
            "Editora web y periodista egresada de Hospital Nubaria. "
            "Estudiante en formación Hospital Nubaria."
        )
        temas, subs = app.etiquetar_sin_llm(
            ["Notas del día en Cali"], [ctx], marca, None
        )
        self.assertEqual(
            app.string_norm_label(temas[0]),
            app.string_norm_label("Estudiantes y egresados"),
        )
        self.assertTrue(app._es_subtema_autoria(subs[0]), subs[0])


class TestCalidadNoRegresa(unittest.TestCase):
    """Known-good Diporto-style phrases must still be produced."""

    def test_proyecto_residencial_de_lujo_se_conserva(self):
        for marca, aliases in (("Diporto", None), ("Serena del Mar", None)):
            with self.subTest(marca=marca):
                sub = app._extraer_subtema_especifico(CTX_DIPORTO, marca, aliases)
                n = app.unidecode(sub.lower())
                self.assertFalse(app._subtema_de_baja_calidad(sub, CTX_DIPORTO), sub)
                self.assertTrue(
                    ("residencial" in n and ("lujo" in n or "serena" in n or "proyecto" in n))
                    or ("proyecto" in n and ("residencial" in n or "lujo" in n)),
                    sub,
                )
                self.assertGreaterEqual(len(sub.split()), 4, sub)

    def test_lanzamiento_carrera_sigue_siendo_valido(self):
        self.assertTrue(app._validar_estructura_subtema("Lanzamiento de carrera deportiva"))
        self.assertTrue(app._candidato_subtema_ok(
            "Lanzamiento de carrera deportiva",
            "La UTB lanza una carrera de medicina deportiva.",
            MARCA, ALIAS,
        ))


class _PBarSpy:
    def __init__(self):
        self.msgs = []

    def progress(self, *args, **kwargs):
        msg = kwargs.get("text")
        if msg is None and len(args) > 1:
            msg = args[1]
        self.msgs.append(msg or "")


class TestTemasGeneralesSinTope(unittest.TestCase):
    """Temas más generales que sus subtemas; acreditación ≠ estudiantes/egresados.

    Sin recorte a 18–30 / Máx:25: mejor N temas coherentes que 25 cubos mezclados.
    """

    _DOMINIOS = [
        ("Acreditación de alta calidad institucional",
         "La universidad renovó su acreditación de alta calidad ante el ministerio."),
        ("Estudiante en formación universitaria",
         "Editor web y periodista egresado de la Universidad Autónoma de Occidente. "
         "Estudiante en formación Universidad Autónoma de Occidente."),
        ("Proyecto residencial de lujo",
         "Diporto propone una experiencia residencial de lujo en Serena del Mar."),
        ("Lanzamiento de carrera deportiva",
         "La universidad lanza una nueva carrera de medicina deportiva en Cartagena."),
        ("Investigación por fallas operativas",
         "La institución enfrenta una investigación por fallas operativas en laboratorios."),
        ("Convenio de formación profesional",
         "La universidad firma un convenio de formación profesional con el SENA."),
        ("Inauguración de laboratorio marino",
         "Inauguraron el laboratorio de biotecnología marina en el campus costero."),
        ("Protesta estudiantil por matrículas",
         "Estudiantes marcharon en protesta por el alza de matrículas."),
        ("Foro de innovación tecnológica",
         "El foro de innovación tecnológica reunió a centros de investigación."),
        ("Campaña de vacunación comunitaria",
         "El hospital lideró una campaña de vacunación comunitaria en el valle."),
        ("Inversión en infraestructura vial",
         "Anunciaron inversión en infraestructura vial para el corredor regional."),
        ("Premio de innovación académica",
         "Recibieron un premio de innovación académica por el laboratorio."),
        ("Exportación del sector avícola",
         "El encuentro avícola analizó las oportunidades de exportación."),
        ("Reforma de política ambiental",
         "Presentaron una reforma de política ambiental para el río."),
        ("Cierre de planta industrial",
         "La empresa anunció el cierre de planta industrial en Buga."),
        ("Alianza de cooperación científica",
         "Suscribieron una alianza de cooperación científica con el SENA."),
        ("Ranking de calidad hospitalaria",
         "El hospital subió en el ranking de calidad hospitalaria nacional."),
        ("Festival de cultura regional",
         "Organizan un festival de cultura regional en el centro de Cali."),
        ("Contratos de empleo temporal",
         "Abrieron contratos de empleo temporal para egresados técnicos."),
        ("Crisis energética del Caribe",
         "Analizan la crisis energética del Caribe y las tarifas."),
        ("Demanda por contaminación hídrica",
         "Presentaron una demanda por contaminación hídrica del río."),
        ("Apertura de sede universitaria",
         "La universidad inauguró una nueva sede universitaria en Palmira."),
        ("Programa de becas rurales",
         "Lanzaron un programa de becas rurales para jóvenes del Pacífico."),
        ("Auditoría de contratación pública",
         "La contraloría abrió una auditoría de contratación pública."),
        ("Cumbre de sostenibilidad ambiental",
         "La cumbre de sostenibilidad ambiental reunió a gobernadores."),
        ("Plataforma de trámites digitales",
         "Estrenaron una plataforma de trámites digitales para matrículas."),
        ("Huelga de personal médico",
         "Hubo huelga de personal médico por el atraso de salarios."),
        ("Ampliación del puerto marítimo",
         "Aprobaron la ampliación del puerto marítimo de Buenaventura."),
    ]

    def test_tema_mas_general_que_subtema_y_nucleos_separados(self):
        titulos, resumenes = [], []
        for i, (tit, res) in enumerate(self._DOMINIOS):
            titulos.append(f"{tit} ({i})")
            resumenes.append(res)
        with patch.object(app, "get_embeddings_batch", return_value=[None] * len(titulos)):
            temas, subs = app.etiquetar_sin_llm(titulos, resumenes, MARCA_UAO, ALIAS_UAO)
            clustered = app.consolidar_temas(subs, resumenes, app._PBarNulo(), MARCA_UAO)
        for tema, sub in zip(clustered, subs):
            self.assertNotEqual(
                app.string_norm_label(tema),
                app.string_norm_label(sub),
                (tema, sub),
            )
        idx_acred = 0
        idx_egres = 1
        self.assertNotEqual(
            app.string_norm_label(clustered[idx_acred]),
            app.string_norm_label(clustered[idx_egres]),
            (clustered[idx_acred], clustered[idx_egres], subs[idx_acred], subs[idx_egres]),
        )
        self.assertFalse(
            app._temas_nucleo_incompatible(clustered[idx_acred], clustered[idx_acred])
        )
        self.assertTrue(
            app._temas_nucleo_incompatible(
                "Acreditación institucional", "Estudiantes y egresados"
            )
        )

    def test_consolidar_300_subtemas_pocos_segundos_sin_embeddings(self):
        import time
        n = 300
        subs, textos = [], []
        for i in range(n):
            tit, res = self._DOMINIOS[i % len(self._DOMINIOS)]
            subs.append(f"{tit} caso {i}")
            textos.append(res)
        spy = _PBarSpy()
        t0 = time.perf_counter()
        with patch.object(app, "get_embeddings_batch", side_effect=AssertionError("no extra embeddings")):
            out = app.consolidar_temas(subs, textos, spy, MARCA_UAO)
        elapsed = time.perf_counter() - t0
        self.assertEqual(len(out), n)
        self.assertLess(elapsed, 5.0, elapsed)
        self.assertTrue(
            any(re.search(r"Temas \d+/\d+", str(m)) for m in spy.msgs),
            spy.msgs[:8],
        )
        self.assertNotIn("Máx:", " ".join(str(m) for m in spy.msgs))
        idx_acred = 0
        idx_egres = 1
        self.assertNotEqual(
            app.string_norm_label(out[idx_acred]),
            app.string_norm_label(out[idx_egres]),
            (out[idx_acred], out[idx_egres], subs[idx_acred], subs[idx_egres]),
        )
        for tema, sub in zip(out, subs):
            self.assertNotEqual(app.string_norm_label(tema), app.string_norm_label(sub), (tema, sub))


class TestXlsxSinColgarTrasTemas(unittest.TestCase):
    """Tras Temas listos: last_grupos, captions que se mueven, Excel sin re-audit."""

    def test_consistencia_400_last_grupos_menos_1s_sin_matcher_all_pairs(self):
        import time as _time
        from difflib import SequenceMatcher
        pd = __import__("pandas")
        n = 400
        titulos = [f"Zeta{i} anuncia hecho puntual {i} en Cali" for i in range(n)]
        resumenes = [f"Resumen corto del hecho {i} en la jornada regional." for i in range(n)]
        df = pd.DataFrame({
            "Título": titulos,
            "Resumen - Aclaracion": resumenes,
            "Tono IA": ["Neutro"] * n,
            "Tema": ["Hecho puntual"] * n,
            "Subtema": [f"Hecho puntual {i}" for i in range(n)],
        })
        last_grupos = {i: [i, i + 1] for i in range(0, n, 2)}
        orig_ratio = SequenceMatcher.ratio
        calls = {"n": 0}

        def spy_ratio(self):
            calls["n"] += 1
            return orig_ratio(self)

        t0 = _time.perf_counter()
        with patch.object(SequenceMatcher, "ratio", spy_ratio), \
             patch.object(app, "construir_grafo_equivalencia",
                          side_effect=AssertionError("no rebuild graph after Temas")):
            out = app.aplicar_consistencia_grupos(
                df, "Título", "Resumen - Aclaracion",
                marca="ZetaCorp", grupos_noticia=last_grupos, saltar_grafo=True,
            )
        elapsed = _time.perf_counter() - t0
        nxn = n * (n - 1) // 2
        self.assertEqual(len(out), n)
        self.assertLess(elapsed, 1.0, elapsed)
        self.assertLess(calls["n"], n, calls["n"])
        self.assertLess(calls["n"], nxn // 10, calls["n"])
        self.assertTrue(all(str(g).startswith("G") for g in out["Grupo noticia"]))
        self.assertEqual(out.loc[0, "Grupo noticia"], out.loc[1, "Grupo noticia"])
        self.assertNotEqual(out.loc[0, "Grupo noticia"], out.loc[2, "Grupo noticia"])

    def test_consistencia_mismo_subtema_une_sin_grafo(self):
        pd = __import__("pandas")
        df = pd.DataFrame({
            "Título": ["Alpha uno", "Beta dos"],
            "Resumen - Aclaracion": ["Resumen A", "Resumen B"],
            "Tono IA": ["Positivo", "Neutro"],
            "Tema": ["Tema A", "Tema B"],
            "Subtema": ["Lanzamiento de carrera deportiva", "Lanzamiento de carrera deportiva"],
        })
        with patch.object(app, "construir_grafo_equivalencia",
                          side_effect=AssertionError("no graph")):
            out = app.aplicar_consistencia_grupos(
                df, "Título", "Resumen - Aclaracion",
                marca="ZetaCorp", grupos_noticia={}, saltar_grafo=True,
            )
        self.assertEqual(out.loc[0, "Grupo noticia"], out.loc[1, "Grupo noticia"])
        self.assertEqual(out.loc[0, "Tono IA"], out.loc[1, "Tono IA"])
        self.assertEqual(out.loc[0, "Tema"], out.loc[1, "Tema"])
        self.assertEqual(out.loc[0, "Subtema"], out.loc[1, "Subtema"])

    def test_generate_output_excel_no_repite_brand_audit_si_hay_contexto(self):
        rows = [{
            "ID Noticia": 1,
            "Título": "UTB lanza carrera de medicina deportiva en Cartagena",
            "Resumen - Aclaracion": "La Universidad Tecnológica de Bolívar abre medicina deportiva.",
            "Contexto analizado": "UTB lanza carrera. La UTB abre medicina deportiva.",
            "Coincidencia marca": "UTB",
            "Origen coincidencia": "Título",
            "Tono IA": "Positivo",
            "Tema": "Educación superior",
            "Subtema": "Lanzamiento de carrera deportiva",
            "Grupo noticia": "G00001",
            "Link Nota": {"value": "Link", "url": "https://example.com/nota"},
        }]
        km = {"titulo": "Título", "resumen": "Resumen - Aclaracion"}
        with patch.object(app, "_brand_audit", side_effect=AssertionError("no re-audit")) as mocked:
            data = app.generate_output_excel(rows, km)
        mocked.assert_not_called()
        self.assertTrue(data and len(data) > 100)
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        self.assertIn("Contexto analizado", headers)
        ctx_i = headers.index("Contexto analizado") + 1
        self.assertIn("UTB", str(ws.cell(2, ctx_i).value))
        link_i = headers.index("Link Nota") + 1
        self.assertTrue(ws.cell(2, link_i).hyperlink)

    def test_temas_listos_no_es_ultimo_tick(self):
        pd = __import__("pandas")
        spy = _PBarSpy()
        n = 12
        subs = [f"Hecho puntual {i}" for i in range(n)]
        textos = [f"Resumen del hecho {i} en la jornada." for i in range(n)]
        app.consolidar_temas(subs, textos, spy, "ZetaCorp")
        df = pd.DataFrame({
            "Título": [f"Titulo {i}" for i in range(n)],
            "Resumen - Aclaracion": textos,
            "Tono IA": ["Neutro"] * n,
            "Tema": ["Hecho institucional"] * n,
            "Subtema": subs,
            "Contexto analizado": textos,
        })
        last_grupos = {i: [i] for i in range(n)}
        app.aplicar_consistencia_grupos(
            df, "Título", "Resumen - Aclaracion",
            marca="ZetaCorp", grupos_noticia=last_grupos, pbar=spy, saltar_grafo=True,
        )
        rows = df.to_dict("records")
        km = {"titulo": "Título", "resumen": "Resumen - Aclaracion"}
        with patch.object(app, "_brand_audit", side_effect=AssertionError("no re-audit")):
            app.generate_output_excel(rows, km, pbar=spy)
        msgs = [str(m) for m in spy.msgs]
        self.assertTrue(any("Temas listos" in m for m in msgs), msgs)
        i_temas = max(i for i, m in enumerate(msgs) if "Temas listos" in m)
        after = msgs[i_temas + 1:]
        self.assertTrue(any("Agrupación" in m for m in after), after)
        self.assertTrue(any("Excel" in m for m in after), after)
        self.assertNotIn("Temas listos", msgs[-1])
        self.assertNotIn("Máx:", " ".join(msgs))

    def test_overlay_y_streamlit_parchean_post_temas(self):
        from pathlib import Path
        overlay = Path(app.__file__).with_name("calidad_etiquetas.py").read_text(encoding="utf-8")
        streamlit = Path(app.__file__).read_text(encoding="utf-8")
        colab = Path(app.__file__).with_name("Grill_API_Colab.txt").read_text(encoding="utf-8")
        self.assertIn("def generate_output_excel", overlay)
        self.assertIn("saltar_grafo", overlay)
        self.assertIn("_dsu_desde_last_grupos_y_subtema", overlay)
        self.assertIn("Escribiendo Excel", overlay)
        self.assertIn("saltar_grafo=True", streamlit)
        self.assertIn("pbar=pb", streamlit)
        self.assertIn("Escribiendo Excel", streamlit)
        self.assertIn("saltar_grafo=True", colab)
        self.assertIn("Escribiendo Excel", colab)
        self.assertNotIn("Máx:25", overlay)
        self.assertNotIn("Máx:25", streamlit.split("exec(")[0])


class TestTituloVideoSinCambio(unittest.TestCase):
    """Título with 'Video | …' stays unchanged in output."""

    def test_video_pipe_se_conserva(self):
        raw = "Video | UTB lanza carrera de medicina deportiva en Cartagena"
        self.assertEqual(app.clean_title_for_output(raw), raw)
        self.assertIn("Video |", app.clean_title_for_output(raw))
        # Grouping still uses the real headline, not 'Video'
        norm = app.normalize_title_for_comparison(raw)
        self.assertIn("carrera", norm)
        self.assertIn("video", norm)


if __name__ == "__main__":
    unittest.main()
